# -*- coding: utf-8 -*-
"""
Exporta Docs/GDD.md para Docs/GDD.docx.

O markdown continua sendo a fonte da verdade (versionada no git).
O .docx e um artefato de leitura/anotacao — edicoes feitas nele nao voltam.

Uso:
    python Docs/export-docx.py                # Docs/GDD.md -> Docs/GDD.docx
    python Docs/export-docx.py entrada.md saida.docx

Requer: pip install python-docx
"""
import io
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# --- aparencia -------------------------------------------------------------
BODY_FONT = "Calibri"
MONO_FONT = "Consolas"          # tem box-drawing; preserva os diagramas ASCII
MONO_SIZE = Pt(8)               # a linha de codigo mais larga tem 94 colunas
CODE_BG = "F2F2F2"
RULE_COLOR = "BFBFBF"
ACCENT = RGBColor(0xC0, 0x1B, 0x3C)

INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)")
LINK = re.compile(r"\[([^\]]*)\]\([^)]*\)")


def shade(el, fill):
    """Pinta o fundo de um paragrafo (w:p) ou de uma celula de tabela (w:tc)."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    props = el.get_or_add_tcPr() if el.tag.endswith("}tc") else el.get_or_add_pPr()
    props.append(shd)


def keep_together(p):
    pPr = p._p.get_or_add_pPr()
    for tag in ("w:keepNext", "w:keepLines"):
        pPr.append(OxmlElement(tag))


def add_runs(p, text, mono=False, bold=False):
    """Escreve `text` em `p` interpretando **negrito**, *italico* e `codigo`."""
    text = LINK.sub(r"\1", text)
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("`") and tok.endswith("`") and len(tok) > 1:
            r = p.add_run(tok[1:-1])
            r.font.name = MONO_FONT
            r.font.size = Pt(9)
        elif tok.startswith("**") and tok.endswith("**"):
            r = p.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith("*") and tok.endswith("*"):
            r = p.add_run(tok[1:-1])
            r.italic = True
        else:
            r = p.add_run(tok)
        if mono:
            r.font.name = MONO_FONT
            r.font.size = MONO_SIZE
        if bold:
            r.bold = True


def add_rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(10)
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), RULE_COLOR)
    pbdr.append(bottom)
    p._p.get_or_add_pPr().append(pbdr)


def add_code(doc, lines):
    """Um unico paragrafo com quebras manuais: mantem o bloco ASCII alinhado."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    pf.left_indent = Inches(0.12)
    pf.line_spacing = 1.0
    keep_together(p)
    shade(p._p, CODE_BG)
    r = p.add_run("\n".join(lines))
    r.font.name = MONO_FONT
    r.font.size = MONO_SIZE
    # garante a fonte tambem para os glifos nao-latinos (box drawing)
    rPr = r._element.get_or_add_rPr().find(qn("w:rFonts"))
    if rPr is not None:
        rPr.set(qn("w:eastAsia"), MONO_FONT)
        rPr.set(qn("w:cs"), MONO_FONT)


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for j, cell in enumerate(header):
        c = t.cell(0, j)
        c.text = ""
        shade(c._tc, "EDEDED")
        add_runs(c.paragraphs[0], cell, bold=True)
    for i, row in enumerate(body, start=1):
        for j in range(len(header)):
            c = t.cell(i, j)
            c.text = ""
            add_runs(c.paragraphs[0], row[j] if j < len(row) else "")
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def convert(md_path, docx_path):
    lines = io.open(md_path, encoding="utf-8").read().split("\n")

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = BODY_FONT
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(6)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.9)
        s.top_margin = s.bottom_margin = Inches(0.8)
    for name, size in (("Heading 1", 20), ("Heading 2", 15), ("Heading 3", 12.5), ("Heading 4", 11)):
        h = doc.styles[name]
        h.font.name = BODY_FONT
        h.font.size = Pt(size)
        h.font.color.rgb = ACCENT if name in ("Heading 1", "Heading 2") else RGBColor(0x33, 0x33, 0x33)

    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            while buf and not buf[-1].strip():
                buf.pop()
            if buf:
                add_code(doc, buf)
            continue

        if stripped.startswith("|") and i + 1 < n and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            rows = [split_row(line)]
            i += 2
            while i < n and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", stripped):
            add_rule(doc)
            i += 1
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            lvl = min(len(m.group(1)), 4)
            p = doc.add_heading("", level=lvl)
            p.paragraph_format.space_before = Pt(14 if lvl <= 2 else 10)
            p.paragraph_format.space_after = Pt(4)
            keep_together(p)
            add_runs(p, m.group(2))
            i += 1
            continue

        if stripped.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(4)
            shade(p._p, "F7F7F7")
            for k, b in enumerate(buf):
                if k:
                    p.add_run("\n")
                add_runs(p, b)
            for r in p.runs:
                r.italic = True
            continue

        m = re.match(r"^(\s*)([-*+]|\d+\.)\s+(.*)$", line)
        if m:
            depth = len(m.group(1)) // 4
            ordered = m.group(2)[0].isdigit()
            style = "List Number" if ordered else "List Bullet"
            if depth:
                style += " %d" % min(depth + 1, 3)
            try:
                p = doc.add_paragraph(style=style)
            except KeyError:
                p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, m.group(3))
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_runs(p, stripped)
        i += 1

    doc.save(docx_path)
    return docx_path


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    src = sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "GDD.md")
    dst = sys.argv[2] if len(sys.argv) > 2 else os.path.join(here, "GDD.docx")
    out = convert(src, dst)
    print("OK -> %s (%.0f KB)" % (out, os.path.getsize(out) / 1024.0))
